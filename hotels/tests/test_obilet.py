"""Offline tests for obilet.py and segments.py (no network).

    py -m unittest discover -s tests -v      (from the hotels folder)
"""

import csv
import gzip
import json
import sys
import tempfile
import unittest
from datetime import date, timedelta
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent))

import obilet  # noqa: E402
import segments  # noqa: E402

FIXTURE = json.loads((HERE / "fixtures" / "obilet_kemer_2ad_1n.json").read_text(encoding="utf-8"))
REGION = dict(region_id="test_kemer", center=dict(lat=36.60279, lon=30.55975), radius_km=5,
              slug="kemer-609-60689_23501_125560-2",
              location=dict(id=609, externalId="60689_23501_125560", locationId=609, name="Kemer",
                            type="2", isCityCenter=False, level=5))
STAY = dict(check_in_date="2026-11-02", check_out_date="2026-11-03", nights=1, kind="night", adults=2)
DAY = date(2026, 10, 28)


class PlanTests(unittest.TestCase):
    CFG = dict(horizon_days=60, start_offset_days=1, adults=[1, 2],
               validation=dict(check_in_offsets=[7, 21, 45], nights=[3, 7], adults=[2]))

    def test_every_night_for_every_adult_count(self):
        stays = obilet.plan_stays(DAY, self.CFG)
        nights = [s for s in stays if s["kind"] == "night"]
        self.assertEqual(len(nights), 59 * 2)  # check-ins D+1 .. D+59, check-out <= D+60
        self.assertEqual(nights[0]["check_in_date"], "2026-10-29")
        self.assertEqual(nights[-1]["check_out_date"], "2026-12-27")

    def test_one_adult_nights_can_be_limited_to_the_first_days(self):
        cfg = {**self.CFG, "adults_horizon_days": {"1": 30}}
        nights = [s for s in obilet.plan_stays(DAY, cfg) if s["kind"] == "night"]
        one = [s for s in nights if s["adults"] == 1]
        self.assertEqual(len(one), 30)
        self.assertEqual(one[-1]["check_in_date"], "2026-11-27")  # day + 30
        self.assertEqual(len(nights) - len(one), 59)

    def test_validation_check_in_rotates_with_the_collection_day(self):
        val = dict(adults=[2], stays=[dict(nights=3, offsets=[3, 4, 5, 6, 7, 8, 9])])
        tuesdays = [date(2026, 10, 27) + timedelta(days=7 * k) for k in range(7)]
        weekdays = {(d + timedelta(days=obilet.validation_stays(d, val)[0]["offset"])).weekday() for d in tuesdays}
        self.assertEqual(len(weekdays), 7)  # seven Tuesday runs check in on all seven weekdays

    def test_real_config_matches_the_flight_window(self):
        cfg = obilet.load_config()
        stays = obilet.plan_stays(date(2026, 10, 30), cfg)
        self.assertTrue(all(s["check_out_date"] <= "2026-11-29" for s in stays))  # day + 30
        self.assertEqual(sum(s["kind"] == "night" for s in stays), 29 * 2)
        self.assertEqual(sorted({s["nights"] for s in stays if s["kind"] == "validation"}), [3, 7, 14])
        self.assertGreaterEqual(cfg["gap_seconds"], 30)

    def test_validation_stays_fit_the_horizon(self):
        val = [s for s in obilet.plan_stays(DAY, self.CFG) if s["kind"] == "validation"]
        self.assertEqual(len(val), 6)
        self.assertTrue(all(s["check_out_date"] <= "2026-12-27" for s in val))
        self.assertTrue(all(s["adults"] == 2 for s in val))

    def test_query_id_depends_on_guests_and_dates(self):
        a = obilet.query_id("antalya", STAY, 1)
        self.assertNotEqual(a, obilet.query_id("antalya", {**STAY, "adults": 1}, 1))
        self.assertNotEqual(a, obilet.query_id("antalya", {**STAY, "check_out_date": "2026-11-04"}, 1))
        self.assertEqual(a, obilet.query_id("antalya", dict(STAY), 1))

    def test_paths(self):
        self.assertEqual(obilet.api_path(REGION, STAY), "/json/oteller/60689_23501_125560/20261102-20261103/2ad/2")
        self.assertIn("/oteller/kemer-609-60689_23501_125560-2/20261102-20261103/2ad?locationLevel=5",
                      obilet.page_url(REGION, STAY))


class ParseTests(unittest.TestCase):
    def setUp(self):
        self.rows, self.facts, self.stats, self.cached = obilet.parse_response(
            FIXTURE, REGION, STAY, 1, DAY, "2026-10-28T12:00:00+03:00")

    def test_radius_filter_drops_far_hotels(self):
        self.assertEqual(self.stats["hotels_total"], 6)
        self.assertEqual(self.stats["hotels_in_radius"], 4)
        self.assertTrue(all(f["distance_to_region_center_km"] <= 5 for f in self.facts.values()))

    def test_price_is_integer_kurus_stay_total(self):
        lancora = next(r for r in self.rows if r["hotel_name"] == "Lancora Beach Hotel")
        self.assertEqual(lancora["price_total_minor"], 605500)
        self.assertEqual(lancora["price_per_night_minor"], 605500)
        self.assertEqual(lancora["board_category_name"], "Her Şey Dahil")
        self.assertTrue(lancora["analysis_eligible"])

    def test_facilities_become_named_columns(self):
        lancora = next(f for f in self.facts.values() if f["hotel_name"] == "Lancora Beach Hotel")
        self.assertEqual(lancora["hotel_type_name"], "Tatil Otelleri")
        self.assertIn("Havuz", lancora["facilities"].split(";"))
        self.assertTrue(lancora["has_pool"])
        self.assertEqual({k for k in obilet.FACILITY_FLAGS} - set(lancora), set())

    def test_no_discount_column_when_the_price_is_not_reduced(self):
        lancora = next(r for r in self.rows if r["hotel_name"] == "Lancora Beach Hotel")
        self.assertIsNone(lancora["price_before_discount_minor"])

    def test_unavailable_hotels_are_kept_without_price(self):
        un = [r for r in self.rows if r["availability_status"] == "unavailable"]
        self.assertTrue(un)
        self.assertTrue(all(r.get("price_total_minor") is None and r["analysis_eligible"] is False for r in un))

    def test_money_unknown_is_none_never_zero(self):
        self.assertIsNone(obilet.to_minor(None))
        self.assertIsNone(obilet.to_minor(0))
        self.assertEqual(obilet.to_minor(1954.0), 195400)
        self.assertEqual(obilet.to_minor("12.345"), 1234)

    def test_dates_echo(self):
        self.assertTrue(obilet.response_matches(FIXTURE, STAY))
        self.assertFalse(obilet.response_matches(FIXTURE, {**STAY, "check_in_date": "2026-11-05"}))

    def test_written_files_round_trip(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "offers.csv.gz"
            obilet.write_csv(path, obilet.OFFER_FIELDS, self.rows, compress=True)
            with gzip.open(path, "rt", encoding="utf-8") as f:
                back = list(csv.DictReader(f))
            self.assertEqual(len(back), len(self.rows))
            self.assertEqual(list(back[0]), obilet.OFFER_FIELDS)


def offer(hid, ci, co, nights, total, board="Kahvaltı Dahil", adults=2, kind="night"):
    return dict(hotel_id=hid, check_in_date=ci, check_out_date=co, nights=str(nights), stay_kind=kind,
                search_adults=str(adults), price_total_minor=str(total) if total else "",
                analysis_eligible="true" if total else "false", board_category_name=board, is_refundable="false")


def hotel(hid, score=8.5, comments=200, stars=4):
    return dict(hotel_id=hid, hotel_name=f"Hotel {hid}", average_score=str(score), comment_count=str(comments),
                stars=str(stars))


class SegmentTests(unittest.TestCase):
    NIGHTS = ["2026-11-14", "2026-11-15", "2026-11-16"]

    def nights_for(self, hid, price, skip=None, board="Kahvaltı Dahil"):
        out = []
        for i, d in enumerate(self.NIGHTS):
            nxt = ["2026-11-15", "2026-11-16", "2026-11-17"][i]
            out.append(offer(hid, d, nxt, 1, None if d == skip else price, board))
        return out

    def test_sum_of_nights_and_terciles(self):
        offers, hotels = [], {}
        for i, price in enumerate([100000, 120000, 150000, 200000, 260000, 400000], 1):
            offers += self.nights_for(str(i), price)
            hotels[str(i)] = hotel(str(i))
        res = segments.tiers_for_stay(offers, hotels, "2026-11-14", "2026-11-17")
        self.assertEqual(res["nights"], 3)
        self.assertEqual([h["hotel_id"] for h in res["tiers"]["economy"]], ["1", "2"])
        self.assertEqual([h["hotel_id"] for h in res["tiers"]["premium"]], ["5", "6"])
        self.assertEqual(res["tiers"]["economy"][0]["stay_total_minor"], 300000)
        self.assertEqual(res["tiers"]["economy"][0]["price_method"], "sum_of_nights")

    def test_hotel_missing_a_night_is_left_out(self):
        offers = self.nights_for("1", 100000, skip="2026-11-15") + self.nights_for("2", 100000)
        res = segments.tiers_for_stay(offers, {"1": hotel("1"), "2": hotel("2")}, "2026-11-14", "2026-11-17")
        self.assertEqual(res["excluded"]["1"], "not bookable on every night")
        self.assertEqual(res["bookable"], 1)

    def test_exact_multi_night_price_wins(self):
        offers = self.nights_for("1", 100000) + [offer("1", "2026-11-14", "2026-11-17", 3, 270000, kind="validation")]
        res = segments.tiers_for_stay(offers, {"1": hotel("1")}, "2026-11-14", "2026-11-17")
        h = res["tiers"]["economy"][0]
        self.assertEqual((h["stay_total_minor"], h["price_method"]), (270000, "exact_stay"))
        self.assertEqual(segments.validation_report(offers)["median_ratio"], round(300000 / 270000, 3))

    def test_quality_meal_and_adults_filters(self):
        offers = (self.nights_for("1", 100000) + self.nights_for("2", 100000, board="Sadece Oda")
                  + self.nights_for("3", 100000))
        hotels = {"1": hotel("1"), "2": hotel("2"), "3": hotel("3", score=6.0)}
        res = segments.tiers_for_stay(offers, hotels, "2026-11-14", "2026-11-17", meal="breakfast")
        self.assertEqual(res["bookable"], 1)
        self.assertEqual(res["excluded"]["3"], "guest score below minimum")
        self.assertIn("meal plan", res["excluded"]["2"])
        self.assertEqual(segments.tiers_for_stay(offers, hotels, "2026-11-14", "2026-11-17", adults=1)["bookable"], 0)


# ---------------- refresh cycles, resume after a refusal (fake site, no network) ----------------

class FakeSite:
    """Answers like obilet for any dates, refuses (403) after `allow` queries."""
    def __init__(self, allow):
        self.allow, self.calls = allow, 0

    def answer(self, path):
        self.calls += 1
        if self.calls > self.allow:
            return dict(status=403, ms=100, text='"Forbidden"')
        a, b = path.split("/")[4].split("-")
        d = json.loads(json.dumps(FIXTURE))
        d["data"]["checkInDate"] = f"{a[:4]}-{a[4:6]}-{a[6:]}T00:00:00"
        d["data"]["checkOutDate"] = f"{b[:4]}-{b[4:6]}-{b[6:]}T00:00:00"
        return dict(status=200, ms=500, text=json.dumps(d))


class FakePage:
    def __init__(self, site):
        self.site = site

    def goto(self, *a, **k):
        pass

    def wait_for_timeout(self, *a):
        pass

    def evaluate(self, js, args):
        return self.site.answer(args[0])


class FakePlaywright:
    def __init__(self, site):
        site_ = site

        class Ctx:
            def new_page(self):
                return FakePage(site_)

            def close(self):
                pass

        class Browser:
            def new_context(self, **k):
                return Ctx()

            def close(self):
                pass

        class Chromium:
            def launch(self, **k):
                return Browser()

        self.chromium = Chromium()


class CycleTests(unittest.TestCase):
    CFG = dict(horizon_days=4, start_offset_days=1, adults=[1, 2], rooms=1, cycle_weekdays=[1, 4],
               validation=dict(adults=[2], stays=[dict(nights=2, offsets=[2])]),
               gap_seconds=0, session_minutes=8, seconds_between_sessions=0, max_queries_per_run=0,
               regions=[REGION, {**REGION, "region_id": "test_two"}])

    def test_cycle_starts_on_tuesday_and_friday(self):
        self.assertEqual(obilet.cycle_start(date(2026, 10, 28), self.CFG), date(2026, 10, 27))  # Wed -> Tue
        self.assertEqual(obilet.cycle_start(date(2026, 10, 30), self.CFG), date(2026, 10, 30))  # Fri
        self.assertEqual(obilet.cycle_start(date(2026, 11, 2), self.CFG), date(2026, 10, 30))   # Mon -> Fri

    def test_jobs_interleave_regions_and_drop_past_check_ins(self):
        regions = self.CFG["regions"]
        jobs = obilet.build_jobs(self.CFG, regions, date(2026, 10, 27), date(2026, 10, 28))
        first = [(r["region_id"], s["check_in_date"], s["adults"]) for r, s in jobs[:4]]
        self.assertEqual(first, [("test_kemer", "2026-10-29", 1), ("test_two", "2026-10-29", 1),
                                 ("test_kemer", "2026-10-29", 2), ("test_two", "2026-10-29", 2)])
        self.assertTrue(all(s["check_in_date"] > "2026-10-28" for _, s in jobs))  # 10-28 is past
        self.assertEqual(jobs[-1][1]["kind"], "validation")

    def test_refusal_stops_the_run_and_the_next_run_resumes(self):
        import types
        with tempfile.TemporaryDirectory() as out:
            cfg_path = Path(out) / "cfg.json"
            cfg_path.write_text(json.dumps(self.CFG), encoding="utf-8")
            fake_module = types.ModuleType("playwright.sync_api")
            site = FakeSite(allow=3)

            class Ctx:
                def __enter__(self_inner):
                    return FakePlaywright(site)

                def __exit__(self_inner, *a):
                    return False

            fake_module.sync_playwright = lambda: Ctx()
            old = sys.modules.get("playwright.sync_api")
            sys.modules["playwright.sync_api"] = fake_module
            try:
                args = ["--config", str(cfg_path), "collect", "--out", out, "--date", "2026-10-28"]
                self.assertEqual(obilet.main(args), 3)                 # refused after 3 answers
                folder = Path(out) / "2026-10-27" / "test_kemer"
                self.assertEqual(len(obilet.done_query_ids(folder)), 2)
                m1 = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
                self.assertFalse(m1["complete"])
                site.allow = 10 ** 6                                 # the next night: no refusal
                self.assertEqual(obilet.main(args), 0)
                m2 = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
                self.assertTrue(m2["complete"])
                self.assertEqual(len(m2["runs"]), 2)
                with gzip.open(folder / "offers.csv.gz", "rt", encoding="utf-8") as f:
                    rows = list(csv.DictReader(f))                   # two gzip members read as one
                self.assertEqual(len({(r["query_id"]) for r in rows}),
                                 m2["planned_queries"] - m2["expired_queries"])
                self.assertEqual(m2["expired_queries"], 2)            # check-in 10-28: 1 and 2 adults
                calls = site.calls
                self.assertEqual(obilet.main(args), 0)                 # complete: nothing is sent
                self.assertEqual(site.calls, calls)
            finally:
                if old is not None:
                    sys.modules["playwright.sync_api"] = old
                else:
                    sys.modules.pop("playwright.sync_api", None)


class ExcelExportTests(unittest.TestCase):
    def test_one_sheet_joins_prices_and_hotel_facts_in_try(self):
        try:
            from openpyxl import load_workbook
        except ImportError:
            self.skipTest("openpyxl not installed")
        import export_excel
        rows, facts, _, _ = obilet.parse_response(FIXTURE, REGION, STAY, 1, DAY, "2026-10-28T12:00:00+03:00")
        with tempfile.TemporaryDirectory() as tmp:
            folder = Path(tmp) / "2026-10-27" / "test_kemer"
            cfg = dict(horizon_days=30, adults=[2], validation={}, rooms=1)
            obilet.save_region(folder, cfg, REGION, date(2026, 10, 27), DAY, rows, facts, [], dict(stopped=False))
            out = Path(tmp) / "x.xlsx"
            export_excel.main(["--data", tmp, "--out", str(out)])
            ws = load_workbook(out, read_only=True)["Fiyatlar"]
            data = list(ws.iter_rows(values_only=True))
            head = data[0]
            lancora = next(r for r in data[1:] if r[head.index("Otel")] == "Lancora Beach Hotel")
            self.assertEqual(lancora[head.index("Toplam fiyat (TL)")], 6055)
            self.assertEqual(lancora[head.index("Pansiyon")], "Her Şey Dahil")
            self.assertEqual(lancora[head.index("Havuz")], "Evet")
            self.assertEqual(lancora[head.index("Yıldız")], 4)


if __name__ == "__main__":
    unittest.main()

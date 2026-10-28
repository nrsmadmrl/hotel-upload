"""One Excel file per refresh cycle: every region, every price, hotel facts on the same row.

For people who analyse the data in Excel rather than code. The CSV files stay the source of
truth; this only joins offers.csv.gz with hotels.csv, converts kurus to TRY and uses Turkish
column names.

    py export_excel.py --data ../../hotel-data/obilet                    # newest cycle
    py export_excel.py --data ../../hotel-data/obilet --cycle 2026-10-30 --out otel_2026-10-30.xlsx
    py export_excel.py ... --available-only                              # priced rows only

Sheets: Fiyatlar (one row per hotel x stay x adults), Oteller, Ozet (per region), Sutunlar
(what every column means).
"""

import argparse
import csv
import gzip
import json
import sys
from pathlib import Path

YES, NO = "Evet", "Hayır"


def yn(v):
    return YES if v == "true" else NO if v == "false" else None


def num(v, cast=float):
    try:
        return cast(v) if v not in (None, "") else None
    except ValueError:
        return None


def try_amount(minor):
    m = num(minor, int)
    return round(m / 100, 2) if m else None


STATUS = {"available": "Müsait", "unavailable": "Müsait değil", "listed_no_price": "Fiyatsız"}
KIND = {"night": "Gecelik", "validation": "Kontrol (çok gecelik)"}

# (Excel column, source, description)
PRICE_COLUMNS = [
    ("Bölge", "area_name", "Pilot bölge (uçuş verisiyle aynı bölgeler)"),
    ("Bölge kodu", "region_id", "Uçuş verisiyle birleştirme anahtarı"),
    ("Otel", "hotel_name", "Otelin Obilet'teki adı"),
    ("Otel ID", "hotel_id", "Obilet otel kimliği; Oteller sayfasıyla birleştirme anahtarı"),
    ("Yıldız", "stars", "Yıldız sınıfı (0 = yıldızsız)"),
    ("Puan", "average_score", "Misafir puanı, 0-10"),
    ("Yorum sayısı", "comment_count", "Puanın dayandığı yorum sayısı"),
    ("Otel tipi", "hotel_type_name", "Şehir Otelleri / Tatil Otelleri"),
    ("Giriş", "check_in_date", "Giriş tarihi"),
    ("Çıkış", "check_out_date", "Çıkış tarihi"),
    ("Gece", "nights", "Konaklama gece sayısı"),
    ("Kişi", "search_adults", "Aranan yetişkin sayısı (1 oda, çocuk yok)"),
    ("Durum", "availability_status", "Müsait / Müsait değil (ör. sezon dışı kapalı) / Fiyatsız"),
    ("Toplam fiyat (TL)", "price_total_minor", "Konaklamanın TOPLAM fiyatı: tüm geceler, 1 oda, aranan bütün kişiler"),
    ("Gecelik fiyat (TL)", "price_per_night_minor", "Toplam fiyat / gece sayısı"),
    ("İndirimsiz fiyat (TL)", "price_before_discount_minor", "Sitedeki üstü çizili fiyat; sadece indirim varsa"),
    ("İndirim oranı", "discount_ratio", "Sitenin gösterdiği indirim oranı; sadece indirim varsa"),
    ("Pansiyon", "board_category_name", "Sadece Oda / Kahvaltı Dahil / Yarım Pansiyon / Her Şey Dahil / ..."),
    ("İade edilebilir", "is_refundable", "Teklif iade edilebilir mi"),
    ("Konaklama vergisi dahil", "accommodation_tax_included", "Sitenin bildirdiği şekilde"),
    ("Satıcı", "supplier", "Teklifi satan acente"),
    ("Havuz", "has_pool", "Obilet'te havuz listeleniyor mu (Hayır = listelenmiyor, kesin yok değil)"),
    ("Otopark", "has_parking", "Obilet'te otopark listeleniyor mu"),
    ("Wi-Fi", "has_wifi", "Obilet'te Wi-Fi listeleniyor mu"),
    ("Denize sıfır / plaj", "is_beachfront", "Deniz kıyısı veya plaj listeleniyor mu"),
    ("Çocuk dostu", "is_child_friendly", "Çocuk dostu veya çocuk havuzu listeleniyor mu"),
    ("Merkeze uzaklık (km)", "distance_to_region_center_km", "Bölge merkezine kuş uçuşu uzaklık"),
    ("Enlem", "latitude", "Otel konumu"),
    ("Boylam", "longitude", "Otel konumu"),
    ("Sorgu türü", "stay_kind", "Gecelik: her gece için 1 gecelik fiyat. Kontrol: gerçek 3/7/14 gecelik fiyat"),
    ("Analize uygun", "analysis_eligible", "Evet: fiyatı ve gece sayısı doğrulanmış satır"),
    ("Girişe kalan gün", "lead_days", "Fiyatın görüldüğü gün ile giriş arasındaki gün"),
    ("Toplama günü", "collection_date", "Fiyatın görüldüğü gün"),
    ("Gözlem zamanı", "observed_at", "Fiyatın görüldüğü an (Türkiye saati)"),
    ("Para birimi", "currency", "TRY"),
    ("Sorgu ID", "query_id", "Aynı sorgunun tekrarlarını eşlemek için"),
]
MONEY = {"price_total_minor", "price_per_night_minor", "price_before_discount_minor"}
BOOL = {"is_refundable", "accommodation_tax_included", "has_pool", "has_parking", "has_wifi",
        "is_beachfront", "is_child_friendly", "is_pet_friendly", "is_accessible", "analysis_eligible"}
NUMBER = {"stars": float, "average_score": float, "comment_count": int, "nights": int,
          "search_adults": int, "discount_ratio": float, "distance_to_region_center_km": float,
          "latitude": float, "longitude": float, "lead_days": int}

HOTEL_COLUMNS = [
    ("Bölge kodu", "region_id"), ("Otel ID", "hotel_id"), ("Otel", "hotel_name"), ("Yıldız", "stars"),
    ("Puan", "average_score"), ("Yorum sayısı", "comment_count"), ("Otel tipi", "hotel_type_name"),
    ("Havuz", "has_pool"), ("Otopark", "has_parking"), ("Wi-Fi", "has_wifi"),
    ("Denize sıfır / plaj", "is_beachfront"), ("Çocuk dostu", "is_child_friendly"),
    ("Evcil hayvan dostu", "is_pet_friendly"), ("Engelli dostu", "is_accessible"),
    ("Tüm özellikler", "facilities"), ("Merkeze uzaklık (km)", "distance_to_region_center_km"),
    ("Enlem", "latitude"), ("Boylam", "longitude"), ("Fotoğraf", "photo_url"),
    ("Müsait görüldüğü sorgu", "seen_available"), ("Müsait değil görüldüğü sorgu", "seen_unavailable"),
]


def cell(field, value):
    if field in MONEY:
        return try_amount(value)
    if field in BOOL:
        return yn(value)
    if field == "availability_status":
        return STATUS.get(value, value or None)
    if field == "stay_kind":
        return KIND.get(value, value or None)
    if field == "facilities":
        return (value or "").replace(";", ", ") or None
    if field in NUMBER:
        return num(value, NUMBER[field])
    if field in ("seen_available", "seen_unavailable"):
        return num(value, int)
    return value if value != "" else None


def read_csv(path):
    opener = gzip.open if str(path).endswith(".gz") else open
    with opener(path, "rt", encoding="utf-8", newline="") as f:
        return list(csv.DictReader(f))


def newest_cycle(root):
    cycles = sorted(p.name for p in Path(root).iterdir()
                    if p.is_dir() and any((p / r / "manifest.json").exists() for r in
                                          [c.name for c in p.iterdir() if c.is_dir()]))
    if not cycles:
        raise SystemExit(f"no cycle folders under {root}")
    return cycles[-1]


def build(root, cycle, available_only=False, area_names=None):
    area_names = area_names or {}
    base = Path(root) / cycle
    prices, hotels, summary = [], [], []
    for folder in sorted(p for p in base.iterdir() if p.is_dir() and (p / "offers.csv.gz").exists()):
        rid = folder.name
        facts = {h["hotel_id"]: h for h in read_csv(folder / "hotels.csv")}
        offers = read_csv(folder / "offers.csv.gz")
        manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
        hotels.extend(facts.values())
        n = 0
        for o in offers:
            if available_only and o.get("availability_status") != "available":
                continue
            row = {**facts.get(o["hotel_id"], {}), **{k: v for k, v in o.items() if v != "" or k not in facts}}
            row["area_name"] = area_names.get(rid, rid)
            prices.append(row)
            n += 1
        summary.append(dict(region=area_names.get(rid, rid), region_id=rid, rows=n, hotels=len(facts),
                            available=sum(1 for o in offers if o.get("availability_status") == "available"),
                            collected=manifest.get("collected_queries"), planned=manifest.get("planned_queries"),
                            expired=manifest.get("expired_queries"), complete=manifest.get("complete")))
    prices.sort(key=lambda r: (r.get("region_id", ""), r.get("check_in_date", ""), int(r.get("search_adults") or 0),
                               r.get("hotel_name") or ""))
    return prices, hotels, summary


def write_xlsx(path, cycle, prices, hotels, summary):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook(write_only=True)
    bold = Font(bold=True)

    def header(ws, names):
        from openpyxl.cell import WriteOnlyCell
        row = []
        for n in names:
            c = WriteOnlyCell(ws, value=n)
            c.font = bold
            row.append(c)
        ws.append(row)

    ws = wb.create_sheet("Fiyatlar")
    ws.freeze_panes = "D2"
    header(ws, [c[0] for c in PRICE_COLUMNS])
    for r in prices:
        ws.append([cell(f, r.get(f)) for _, f, _ in PRICE_COLUMNS])

    ws = wb.create_sheet("Oteller")
    ws.freeze_panes = "D2"
    header(ws, [c[0] for c in HOTEL_COLUMNS])
    for h in sorted(hotels, key=lambda h: (h.get("region_id", ""), h.get("hotel_name") or "")):
        ws.append([cell(f, h.get(f)) for _, f in HOTEL_COLUMNS])

    ws = wb.create_sheet("Ozet")
    header(ws, ["Bölge", "Bölge kodu", "Fiyat satırı", "Otel", "Müsait satır", "Toplanan sorgu",
                "Planlanan sorgu", "Süresi geçen sorgu", "Döngü tamam"])
    for s in summary:
        ws.append([s["region"], s["region_id"], s["rows"], s["hotels"], s["available"], s["collected"],
                   s["planned"], s["expired"], YES if s["complete"] else NO])
    ws.append([])
    ws.append([f"Döngü başlangıcı: {cycle}. Kaynak: obilet (liste fiyatları; rezervasyon/ödeme fiyatı değil)."])
    ws.append(["Boş hücre = bilinmiyor (0 değil). Gecelik satırlardan bir konaklamanın fiyatı, gecelerin "
               "toplamıdır; Kontrol satırları bu toplamın gerçek çok gecelik fiyata yakınlığını ölçer."])

    ws = wb.create_sheet("Sutunlar")
    header(ws, ["Sütun", "Açıklama", "CSV alanı"])
    for name, field, desc in PRICE_COLUMNS:
        ws.append([name, desc, field])
    wb.save(path)


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawTextHelpFormatter)
    ap.add_argument("--data", required=True, help="the obilet folder of the data repository")
    ap.add_argument("--cycle", default=None, help="cycle start, e.g. 2026-10-30 (default: newest)")
    ap.add_argument("--out", default=None, help="output .xlsx (default: otel_fiyatlari_<cycle>.xlsx)")
    ap.add_argument("--available-only", action="store_true", help="only priced rows (smaller file)")
    ap.add_argument("--config", default=str(Path(__file__).resolve().parent / "obilet_regions.json"))
    args = ap.parse_args(argv)
    cycle = args.cycle or newest_cycle(args.data)
    cfg = json.loads(Path(args.config).read_text(encoding="utf-8"))
    names = {r["region_id"]: r.get("area_name", r["region_id"]) for r in cfg["regions"]}
    prices, hotels, summary = build(args.data, cycle, args.available_only, names)
    if len(prices) > 1_048_000:
        raise SystemExit(f"{len(prices)} rows exceed Excel's limit; use --available-only")
    out = args.out or f"otel_fiyatlari_{cycle}.xlsx"
    write_xlsx(out, cycle, prices, hotels, summary)
    print(f"{out}: {len(prices)} price rows, {len(hotels)} hotels, {len(summary)} regions (cycle {cycle})")
    return 0


if __name__ == "__main__":
    sys.exit(main())
